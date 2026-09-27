from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = r"G:\\codex test\\周文豪_测试工程师_模拟面试题与参考答案.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Cm(1.55); sec.bottom_margin = Cm(1.35)
sec.left_margin = Cm(1.7); sec.right_margin = Cm(1.7)

BLACK = RGBColor(0, 0, 0)
NAVY = RGBColor(31, 67, 100)
GRAY = RGBColor(90, 98, 108)

def font(run, size=10, bold=False, color=BLACK):
    run.font.name = 'Microsoft YaHei'
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    rpr.rFonts.set(qn('w:ascii'), 'Microsoft YaHei')
    rpr.rFonts.set(qn('w:hAnsi'), 'Microsoft YaHei')
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = color

def para(p, before=0, after=5, line=1.15, keep=False):
    f = p.paragraph_format
    f.space_before = Pt(before); f.space_after = Pt(after); f.line_spacing = line
    if keep:
        pPr = p._p.get_or_add_pPr(); pPr.append(OxmlElement('w:keepNext'))

def add_text(p, text, size=10, bold=False, color=BLACK):
    r = p.add_run(text); font(r, size, bold, color); return r

def heading(text):
    p = doc.add_paragraph(); para(p, before=11, after=5, line=1.0, keep=True)
    add_text(p, text, 12, True, BLACK)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), '5'); bot.set(qn('w:space'), '5'); bot.set(qn('w:color'), 'D6DEE7')
    pbdr.append(bot); pPr.append(pbdr)

def qa(n, q, a):
    p = doc.add_paragraph(); para(p, before=5, after=2, line=1.08, keep=True)
    add_text(p, f'{n}. {q}', 10.2, True, NAVY)
    p = doc.add_paragraph(); para(p, after=4, line=1.18)
    add_text(p, '参考答案：', 9.5, True, BLACK); add_text(p, a, 9.5)

# Title and opening note
p = doc.add_paragraph(); p.style = doc.styles['Title']; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; para(p, after=3, line=1.0)
add_text(p, '软件测试工程师模拟面试题与参考答案', 20, True, BLACK)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; para(p, after=12, line=1.0)
add_text(p, '根据周文豪简历整理｜功能测试 · 接口测试 · 自动化测试 · OTA 测试', 9.5, False, GRAY)
p = doc.add_paragraph(); para(p, after=7, line=1.18)
add_text(p, '使用说明：', 9.5, True, NAVY)
add_text(p, '以下答案基于简历中的真实经历进行组织，面试时应根据实际参与深度、真实项目细节和现场追问进行调整，不要机械背诵。', 9.5)

heading('一 自我介绍与求职动机')
qa(1, '请做一下自我介绍。', '我目前就读于湖南涉外经济学院软件工程专业，具备软件测试、接口测试、UI 自动化和日志分析能力。在 Soundcore 实习期间，我主要负责耳机固件和 App 的功能测试、OTA 自动化验证、稳定性测试和日志分析，累计编写并执行测试用例 300 多条，通过 Jira 跟踪缺陷 40 多个，其中严重及以上缺陷 8 个。基于 Python 和 Pytest 编写 OTA 自动化脚本后，单轮回归时间由约 2 小时缩短到 20 分钟，效率提升约 80%。此外，我还参与过电商商城系统测试，使用 Postman 进行接口测试，使用 MySQL 校验业务数据，并通过 Pytest 和 Playwright 编写 UI 自动化用例。我希望继续在软件测试和自动化测试方向发展。')
qa(2, '你为什么想从事软件测试？', '我对软件质量保障和问题定位比较感兴趣。测试不仅是执行用例，还需要理解业务、分析风险、设计场景，并通过日志和数据定位问题。在 Soundcore 实习期间，我接触了固件、App、蓝牙连接和 OTA 升级等软硬件结合的测试场景，发现自己比较擅长发现异常、整理问题和推动闭环。同时，我也在学习 Python、Pytest、Playwright 和接口自动化，希望未来向自动化测试和测试开发方向发展。')

heading('二 测试理论与测试流程')
qa(3, '你了解完整的软件测试流程吗？', '我理解的软件测试流程主要包括需求分析、测试计划制定、测试用例设计、测试环境准备、测试执行、缺陷提交与跟踪、回归测试以及测试报告输出。需求分析阶段重点关注功能逻辑、异常流程、边界条件和需求一致性。执行时先进行冒烟测试，确认主流程可用后，再开展功能、接口、兼容性和回归测试。发现问题后，在 Jira 中记录复现步骤、实际结果、预期结果、环境信息和相关日志，修复后进行回归验证。')
qa(4, '你通常如何设计测试用例？', '我会结合需求特点选择等价类划分、边界值分析、场景法、判定表和正交实验法。例如测试登录功能时，我会覆盖正确账号密码、错误账号密码、空值、超长输入、特殊字符、连续登录失败、验证码失效和网络异常等场景，同时考虑不同设备、系统版本和接口异常返回。')
qa(5, '你如何确定回归测试范围？', '我会根据本次修改内容、影响模块、历史缺陷和核心业务流程确定回归范围。如果修改的是 OTA 模块，除了验证本次修复的问题，还需要回归升级包校验、正常升级、断点续传、失败回滚、版本校验和升级后的基础功能；如果修改的是订单模块，则重点回归库存、订单状态、支付结果和用户数据一致性。')

heading('三 Soundcore 实习与 OTA 测试')
qa(6, '你在 Soundcore 主要负责哪些测试工作？', '我主要负责耳机固件和 App 的功能测试，覆盖音频播放、蓝牙连接、多设备切换、降噪和通透模式、触控操作、EQ 音效、充电低电以及 OTA 升级等场景。同时参与 Android 和 iOS 跨机型兼容性测试、长时间播放和反复配对等稳定性测试，并使用串口和 ADB 抓取耳机及 App 日志，辅助分析蓝牙断连和音频卡顿等问题。')
qa(7, '请介绍你的 OTA 自动化测试。', '我使用 Python + Pytest 编写 OTA 自动化脚本，主要验证升级包校验、设备连接、升级过程、版本校验、断点续传和失败回滚等场景。测试前确认设备状态和当前固件版本，执行升级后校验升级结果和基础功能。异常场景下保存执行日志，便于定位失败原因。自动化前单轮回归约 2 小时，使用脚本后缩短到约 20 分钟，覆盖 20 多条 OTA 用例，回归效率提升约 80%。')
qa(8, '如果 OTA 期间设备断电并且无法启动，你如何定位？', '我会同时抓取耳机端和手机端日志，并记录升级包版本、设备型号、系统版本、网络环境、升级进度和故障时间点。然后对比正常升级和异常升级日志，重点关注升级包校验、数据传输、分区写入、重启和回滚信息。AI 可以辅助日志聚类和关键字段提取，但我会结合原始日志与复现结果进行验证。最后整理复现步骤、影响范围、环境、关键日志和初步原因，提交 Jira 并在修复后验证正常升级、失败回滚和重新启动。')
qa(9, '你如何使用 AI 辅助日志分析？', '我会先抓取耳机端和手机端日志，记录故障时间点、操作步骤和设备环境，再将日志和问题现象一起提供给 AI，辅助完成日志聚类、关键字段提取和异常时间点定位。但不会完全依赖 AI 的结论，而是回到原始日志和复现结果进行验证。提交缺陷时保留关键日志和验证结果，确保判断可靠。通过这种方式，蓝牙断连和音频卡顿等问题的平均定位时间缩短约 50%。')

heading('四 缺陷管理与质量保障')
qa(10, '请介绍一个你发现的严重缺陷。', '我曾遇到过 OTA 升级失败后设备无法正常连接的问题。首先确认复现条件，并记录设备型号、固件版本、手机型号、系统版本、网络环境和故障时间。随后同时抓取手机端 App 日志和耳机端串口日志，对比正常与异常升级过程，发现异常情况下升级进入固件写入阶段后没有正确触发回滚，设备处于异常状态。我将复现步骤、影响范围、关键日志和预期结果整理后，通过 Jira 提交给开发和固件团队。修复后重新验证正常升级、断网、断电、升级包异常和失败回滚等场景，确认问题闭环。')
qa(11, '一个合格的缺陷报告应该包含哪些内容？', '至少包括缺陷标题、测试环境、前置条件、详细复现步骤、实际结果、预期结果、复现概率、影响范围、严重程度以及截图或日志。耳机和 App 问题还应补充设备型号、手机型号、系统版本、固件版本和网络环境。标题要简洁准确，步骤要让其他人能够稳定复现。')
qa(12, '如果开发认为这不是缺陷，你会怎么处理？', '我会先确认需求、设计文档和产品规则，避免凭个人理解判断，然后提供清晰的复现步骤、环境、实际结果、预期结果和日志证据。如果需求存在歧义，我会邀请产品、开发和测试共同确认最终规则，并将结论同步到缺陷或需求文档中。重点是基于事实确认产品最终应该表现成什么样。')

heading('五 接口 数据库与 UI 自动化')
qa(13, '使用 Postman 做接口测试时主要关注什么？', '我主要关注请求参数、请求方法、响应状态码、响应数据、异常场景和接口之间的数据关联。以登录接口为例，会验证正确登录、错误密码、空参数、参数格式错误、账号不存在和重复登录等场景，并检查响应字段、错误提示和 Token 是否正确。在电商项目中，我使用 Postman 对核心接口进行参数化测试，完成 60 多条用例回归。')
qa(14, '你在测试中如何使用 MySQL？', '我主要使用 MySQL 进行数据准备、结果校验和业务一致性验证。例如测试下单功能时，检查订单状态是否正确生成、商品库存是否扣减、用户信息是否关联正确，以及支付成功或失败后订单状态是否正确变化。在电商项目中，我通过数据库校验库存、订单状态和用户数据一致性，推动修复了 6 处业务逻辑缺陷。')
qa(15, '你使用 Playwright 做过哪些自动化测试？', '我使用 Pytest + Playwright 编写过电商系统下单和支付主流程的 UI 自动化用例，覆盖登录、选择商品、加入购物车、提交订单和支付等步骤。编写时使用稳定的元素定位方式，并增加必要的等待、断言和异常处理，减少页面加载或元素变化导致的误报。项目中编写了 30 多条 UI 自动化用例，将冒烟测试时间由 40 分钟缩短至 8 分钟。')
qa(16, '是不是所有测试都应该自动化？', '不是。自动化更适合频率高、流程稳定、结果容易判断、重复性强的测试，例如冒烟、回归、接口和固定流程测试。探索性测试、视觉体验测试、需求经常变化的功能以及需要大量人工判断的场景，更适合人工测试。自动化的目标是减少重复劳动、提高回归效率和覆盖率，而不是完全替代人工。')

heading('六 兼容性与综合能力')
qa(17, '你如何开展 Android 和 iOS 兼容性测试？', '我会根据用户群体和产品覆盖范围选择不同品牌、不同系统版本和不同性能层级的设备，重点验证核心功能和高风险场景。耳机 App 主要测试蓝牙连接、多设备切换、音频播放、降噪模式、OTA 升级和异常重连，并记录设备型号、系统版本、App 版本、固件版本和复现结果。实习期间执行了 100 多轮跨机型测试，推动 5 类高频问题在量产前闭环。')
qa(18, '你认为自己的优势是什么？', '第一，我有真实的软硬件结合测试经历，接触过耳机固件、App、OTA、蓝牙和日志分析。第二，我重视量化结果，编写并执行过 300 多条用例，跟踪 40 多个缺陷，并将 OTA 回归时间从 2 小时缩短到 20 分钟。第三，我不仅会执行测试，也会关注问题定位和缺陷闭环，能够使用日志、数据库和接口数据辅助分析。')

heading('七 面试结束时可以反问的问题')
for item in [
    '目前团队主要使用哪些测试工具和自动化框架？',
    '这个岗位更侧重功能测试、接口测试还是自动化测试？',
    '新人入职后主要负责哪些业务模块？',
    '团队对测试人员的成长路径和技术提升有哪些安排？',
    '目前团队在自动化测试或质量保障方面最大的挑战是什么？',
]:
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.left_indent = Cm(0.45); p.paragraph_format.first_line_indent = Cm(-0.2); para(p, after=2, line=1.1); add_text(p, item, 9.5)

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER; para(footer, line=1.0)
add_text(footer, '周文豪 · 测试工程师面试准备', 8, False, GRAY)

doc.core_properties.title = '软件测试工程师模拟面试题与参考答案'
doc.core_properties.author = '周文豪'
doc.save(OUT)
print(OUT)

